from __future__ import annotations
import csv
import io
import time
from uuid import UUID
from schemas.tool import ToolCall, ToolResult
from thrall.tools.filesystem._resolve import resolve, is_protected

_MAX_ROWS = 500


def execute(call: ToolCall) -> ToolResult:
    start = time.monotonic()
    path = resolve(call.args.get("path", ""))

    if is_protected(path) or not path.exists():
        return _result(call.id, error=f"path not found: {path}", start=start)
    if not path.is_file():
        return _result(call.id, error=f"not a file: {path}", start=start)

    ext = path.suffix.lower()
    if ext == ".xlsx":
        return _read_excel(call, path, start)

    max_rows = int(call.args.get("max_rows", _MAX_ROWS))
    include_header = call.args.get("include_header", True)

    try:
        raw = path.read_bytes()
        # Detect encoding — try utf-8-sig first (handles BOM), then latin-1 as fallback
        for encoding in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            return _result(call.id, error="could not decode file — unknown encoding", start=start)

        reader = csv.DictReader(io.StringIO(text))
        headers = reader.fieldnames or []
        rows = []
        for i, row in enumerate(reader):
            if i >= max_rows:
                break
            rows.append(row)

        total_rows = i + 1 if rows else 0
        truncated = total_rows >= max_rows

        lines = []
        if include_header and headers:
            lines.append(" | ".join(str(h) for h in headers))
            lines.append("-" * min(80, len(lines[0])))

        for row in rows:
            lines.append(" | ".join(str(row.get(h, "")) for h in headers))

        output = "\n".join(lines)
        if truncated:
            output += f"\n\n[truncated at {max_rows} rows — file may contain more]"

        header = f"[CSV: {path.name} | {len(headers)} columns, {total_rows} rows shown]\n\n"
        return _result(call.id, output=header + output, start=start)

    except PermissionError:
        return _result(call.id, error=f"permission denied: {path}", start=start)
    except Exception as e:
        return _result(call.id, error=str(e), start=start)


def _read_excel(call: ToolCall, path, start: float) -> ToolResult:
    try:
        import openpyxl
    except ImportError:
        return _result(call.id, error="openpyxl is not installed — run: pip install openpyxl", start=start)

    max_rows = int(call.args.get("max_rows", _MAX_ROWS))
    sheet_name = call.args.get("sheet", None)

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name and sheet_name in wb.sheetnames else wb.active
        sheet_label = ws.title if ws else "Sheet1"

        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if header_row is None:
            return _result(call.id, output=f"[XLSX: {path.name} | empty sheet]", start=start)

        headers = [str(c) if c is not None else "" for c in header_row]
        lines = [" | ".join(headers), "-" * min(80, sum(len(h) for h in headers) + 3 * len(headers))]

        count = 0
        truncated = False
        for row in rows_iter:
            if count >= max_rows:
                truncated = True
                break
            lines.append(" | ".join(str(c) if c is not None else "" for c in row))
            count += 1

        output = "\n".join(lines)
        if truncated:
            output += f"\n\n[truncated at {max_rows} rows]"

        sheets_info = f"sheets: {', '.join(wb.sheetnames)}" if len(wb.sheetnames) > 1 else f"sheet: {sheet_label}"
        header = f"[XLSX: {path.name} | {len(headers)} columns, {count} rows shown | {sheets_info}]\n\n"
        return _result(call.id, output=header + output, start=start)

    except PermissionError:
        return _result(call.id, error=f"permission denied: {path}", start=start)
    except Exception as e:
        return _result(call.id, error=str(e), start=start)


def _result(call_id: UUID, start: float, output: str | None = None, error: str | None = None) -> ToolResult:
    return ToolResult(call_id=call_id, output=output, error=error, duration_ms=int((time.monotonic() - start) * 1000))


NAME = "documents_read_csv"
DESCRIPTION = "Read a CSV or Excel (.xlsx) file. Returns headers and rows as pipe-separated text. Supports encoding detection and row limits."
PARAMETERS = {
    "path":           {"type": "string",  "required": True},
    "max_rows":       {"type": "integer", "required": False, "default": 500},
    "include_header": {"type": "boolean", "required": False, "default": True},
    "sheet":          {"type": "string",  "required": False, "description": "Sheet name for .xlsx files (defaults to active sheet)"},
}
